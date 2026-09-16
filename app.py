import calendar, io
from datetime import timedelta
import numpy as np
import pandas as pd
import plotly.express as px
import streamlit as st
from cleaning import clean_uploads, load_data, save_batch, upload_summary, clean_dispatch_upload, save_dispatch_batch, load_dispatch_data, DatabaseConnectionError

st.set_page_config(page_title="Weekly Itemwise Sales", page_icon="Sales", layout="wide")
WEEKDAYS=list(calendar.day_name)


@st.cache_data(show_spinner=False)
def sales_data(): return load_data()
@st.cache_data(show_spinner=False)
def dispatch_data(): return load_dispatch_data()

def download_pattern(pattern, adjustment_percent=10):
    report = pattern.copy()
    for day in WEEKDAYS:
        if day in report:
            report[day] = np.ceil(pd.to_numeric(report[day], errors="coerce") * (1 + adjustment_percent / 100))
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        report.to_excel(writer, sheet_name="Mon-Sun Pattern", index=False, startrow=3)
        ws = writer.sheets["Mon-Sun Pattern"]
        ws["A1"] = "Monday-Sunday Production Pattern"
        ws["A2"] = f"Includes {adjustment_percent}% production increase"
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(report.columns))
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        navy, blue, white = "17365D", "D9EAF7", "FFFFFF"
        ws["A1"].font = Font(bold=True, size=16, color=white)
        ws["A1"].fill = PatternFill("solid", fgColor=navy)
        ws["A1"].alignment = Alignment(horizontal="center")
        ws["A2"].font = Font(italic=True, color="555555")
        header_row = 4
        thin = Side(style="thin", color="B7C9D6")
        for cell in ws[header_row]:
            cell.font = Font(bold=True, color=navy)
            cell.fill = PatternFill("solid", fgColor=blue)
            cell.alignment = Alignment(horizontal="center")
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for row in ws.iter_rows(min_row=header_row + 1, max_row=ws.max_row):
            for cell in row:
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if isinstance(cell.value, float) and pd.isna(cell.value):
                    cell.value = None
                elif isinstance(cell.value, (int, float)):
                    cell.number_format = "0"
        ws.freeze_panes = "B5"
        ws.auto_filter.ref = f"A4:{chr(64 + len(report.columns))}{ws.max_row}"
        ws.column_dimensions["A"].width = 34
        for column in range(2, len(report.columns) + 1):
            ws.column_dimensions[chr(64 + column)].width = 15
    return out.getvalue()
def average_items(data, metric, day_count, label):
    result=data.groupby(["item_name","unit"],as_index=False)[metric].sum()
    result[label]=result[metric]/day_count
    return result[["item_name","unit",label]].sort_values(label,ascending=False)

def sales_page():
    st.title("Weekly Itemwise Sales Dashboard")
    st.caption("All sales and quantity results are averages. Each weekday total is divided by the number of that weekday business dates in the selected period.")
    with st.sidebar:
        st.header("Sales / cancel upload")
        sf=st.file_uploader("Sales file",type=["csv","xlsx","xls"],key="sales")
        cf=st.file_uploader("Cancel file",type=["csv","xlsx","xls"],key="cancel")
        if st.button("Clean and save sales data",type="primary",disabled=not(sf and cf)):
            try:
                with st.spinner("Cleaning, standardising item names and saving..."):
                    cleaned=clean_uploads(sf,cf); bid,saved=save_batch(cleaned,sf,cf)
                st.cache_data.clear(); st.success(f"{'Saved' if saved else 'Already saved'} {len(cleaned):,} cleaned records ({bid}).")
            except Exception as error: st.error(str(error))
    data=sales_data()
    if data.empty:
        st.info("Upload one Sales file and one Cancel file to start."); return
    data["unit"]=data.get("unit","units").fillna("units").replace("","units")
    with st.sidebar:
        st.divider(); st.header("Current dashboard filters")
        lo,hi=data.date.min().date(),data.date.max().date(); mode=st.radio("Period",["Custom dates","A month","Latest 7 days","Previous 7 days"],key="period")
        if mode=="A month":
            months=sorted(data.month.dropna().unique(),key=pd.to_datetime); choice=st.selectbox("Month",months); ds=data.loc[data.month.eq(choice),"date"]; start,end=ds.min().date(),ds.max().date()
        elif mode=="Latest 7 days": start,end=max(lo,hi-timedelta(days=6)),hi
        elif mode=="Previous 7 days": end=hi-timedelta(days=7); start=max(lo,end-timedelta(days=6))
        else: start,end=st.date_input("Date range",value=(lo,hi),min_value=lo,max_value=hi)
        outlets=st.multiselect("Outlet",sorted(data.outlet.dropna().unique())); sources=st.multiselect("Order source",sorted(data.order_source.dropna().unique())); statuses=st.multiselect("Status",["Completed","Cancelled"],default=["Completed"],help="Completed is selected by default so cancellation rows do not make weekly results negative.")
        st.divider(); st.header("Person segregation"); handlers=st.multiselect("Handled by",sorted(data.handler.dropna().unique()))
    scope=data[(data.date.dt.date>=start)&(data.date.dt.date<=end)].copy()
    if outlets: scope=scope[scope.outlet.isin(outlets)]
    if sources: scope=scope[scope.order_source.isin(sources)]
    filtered=scope[scope.status.isin(statuses)].copy()
    if handlers: filtered=filtered[filtered.handler.isin(handlers)]
    if filtered.empty: st.warning("No data matches these filters."); return
    filtered["display_sales"] = pd.to_numeric(filtered["net_sales"], errors="coerce").fillna(0).abs()
    selected=st.radio("Measure",["Quantity","Sales (INR)"],horizontal=True); metric,title=("analysis_quantity","Quantity") if selected=="Quantity" else ("display_sales","Sales (INR)")
    all_dates=scope.date.dt.date.drop_duplicates(); total_days=len(all_dates); completed=filtered[filtered.status.eq("Completed")]; cancelled=filtered[filtered.status.eq("Cancelled")]
    a,b,c,d=st.columns(4); a.metric(f"Average daily {title}",f"{filtered[metric].sum()/total_days:,.2f}"); b.metric(f"Average completed {title}",f"{completed[metric].sum()/total_days:,.2f}"); c.metric("Average cancelled value",f"{cancelled[metric].abs().sum()/total_days:,.2f}"); d.metric("Business days included",f"{total_days:,}")
    st.subheader("Overall average per business day")
    left,right=st.columns(2); daily=filtered.groupby("date",as_index=False)[metric].sum(); average_label=f"Average {title} per day"; top=average_items(filtered,metric,total_days,average_label).head(25)
    with left: st.plotly_chart(px.line(daily,x="date",y=metric,markers=True,title="Daily actual trend",labels={metric:title}),use_container_width=True)
    with right: st.plotly_chart(px.bar(top.sort_values(average_label),x=average_label,y="item_name",color="unit",orientation="h",title="Top 25 items: daily average",labels={"item_name":"Item","unit":"Unit"}),use_container_width=True)
    st.subheader("Item-wise weekday average")
    st.caption("Every selected business-date occurrence of a weekday is included in its average. If an item sold zero on one Sunday, that Sunday is still counted.")
    weekday_dates=scope[["date","weekday"]].drop_duplicates()
    for day,tab in zip(WEEKDAYS,st.tabs(WEEKDAYS)):
        with tab:
            count=weekday_dates.loc[weekday_dates.weekday.eq(day),"date"].nunique(); weekday_data=filtered[filtered.weekday.eq(day)]
            if count==0: st.info(f"No {day} data for these filters.")
            else:
                label=f"Average {title} per {day}"; report=average_items(weekday_data,metric,count,label)
                st.caption(f"Calculated across {count} {day}{'' if count==1 else 's'}.")
                st.plotly_chart(px.bar(report.head(35),x="item_name",y=label,color="unit",title=f"{day}: average {title} by item",labels={"item_name":"Item","unit":"Unit"}).update_xaxes(tickangle=-45),use_container_width=True)
                st.dataframe(report.rename(columns={"item_name":"Item","unit":"Unit"}),hide_index=True,use_container_width=True)
    with st.expander("Stored data and download"):
        st.dataframe(upload_summary(),hide_index=True,use_container_width=True); st.download_button("Download filtered cleaned sales CSV",filtered.to_csv(index=False).encode(),"filtered_cleaned_sales.csv","text/csv")
def production_page():
    st.title("Dispatch & Production Planning")
    st.caption("Upload dispatch data to calculate weekday demand and recommended production from the 75th-percentile weekly requirement.")
    with st.sidebar:
        st.header("Dispatch upload")
        upload=st.file_uploader("Dispatch file",type=["csv","xlsx","xls"],key="dispatch")
        if st.button("Clean and save dispatch data",type="primary",disabled=not upload):
            try:
                clean=clean_dispatch_upload(upload); bid,saved=save_dispatch_batch(clean,upload); st.cache_data.clear(); st.success(f"{'Saved' if saved else 'Already saved'} {len(clean):,} dispatch rows ({bid}).")
            except Exception as e: st.error(str(e))
    data=dispatch_data()
    if data.empty: st.info("Upload a dispatch file with Item Name, Quantity Delivered and Transfer Date."); return
    with st.sidebar:
        st.divider(); st.caption("The downloaded Monday-Sunday pattern includes a fixed 10% production increase.")
        chosen=st.multiselect("Production item",sorted(data.item_name.unique()))
    if chosen: data=data[data.item_name.isin(chosen)]
    daily=data.groupby(["transfer_date","item_name","weekday"],as_index=False).quantity_delivered.sum()
    weekly=daily.assign(week_start=daily.transfer_date-pd.to_timedelta(daily.transfer_date.dt.dayofweek,unit="D")).groupby(["week_start","item_name"],as_index=False).quantity_delivered.sum().rename(columns={"quantity_delivered":"Weekly Quantity"})
    benchmark=weekly.groupby("item_name")["Weekly Quantity"].agg(Average_Weekly="mean",Median_Weekly="median",P75_Weekly=lambda x:x.quantile(.75),P90_Weekly=lambda x:x.quantile(.90),Min_Weekly="min",Max_Weekly="max").reset_index()
    benchmark["Recommended Weekly Production"]=np.ceil(benchmark.P75_Weekly*1.10).astype(int)
    day_benchmark=daily.groupby(["item_name","weekday"])["quantity_delivered"].agg(Average="mean",Median="median",P75=lambda x:x.quantile(.75),P90=lambda x:x.quantile(.90)).reset_index(); day_benchmark.weekday=pd.Categorical(day_benchmark.weekday,categories=WEEKDAYS,ordered=True); day_benchmark=day_benchmark.sort_values(["item_name","weekday"]); matrix=day_benchmark.pivot(index="item_name",columns="weekday",values="Median").reindex(columns=WEEKDAYS).reset_index()
    st.subheader("Recommended weekly production")
    st.dataframe(benchmark.round(0),hide_index=True,use_container_width=True)
    tabs=st.tabs(WEEKDAYS)
    for day,tab in zip(WEEKDAYS,tabs):
        with tab:
            view=day_benchmark[day_benchmark.weekday.eq(day)].sort_values("P75",ascending=False)
            st.plotly_chart(px.bar(view,x="item_name",y="P75",title=f"{day} production benchmark (P75)").update_xaxes(tickangle=-45),use_container_width=True)
            st.dataframe(view.round(2),hide_index=True,use_container_width=True)
    output=download_pattern(matrix, 10)
    st.download_button("Download Monday-Sunday pattern (+10%)",output,"monday_sunday_production_pattern_plus_10.xlsx","application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")

page=st.sidebar.radio("Page",["Sales dashboard","Dispatch & production"])
try:
    if page=="Sales dashboard": sales_page()
    else: production_page()
except DatabaseConnectionError as error:
    st.error(str(error))
    st.info("Your existing local data remains safe. Correct the Streamlit DATABASE_URL Secret, save it, and reboot the app.")



